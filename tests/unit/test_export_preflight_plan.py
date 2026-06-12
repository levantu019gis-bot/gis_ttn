from __future__ import annotations

from datetime import date, time
from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from thucthengay.export import build_export_preflight_plan, ensure_final_renders_for_export
from thucthengay.models import (
    Composition,
    CompositionArtifacts,
    GridConfig,
    GridInterval,
    ImageLayer,
    Issue,
    IssueScope,
    IssueSeverity,
    MapFrame,
    MetadataStatus,
    PlaceholderType,
    TargetConfig,
    TemplateMetadata,
    TemplatePlaceholder,
    ViewState,
)
from thucthengay.render import RasterRenderResult, RenderSpec
from thucthengay.workspace import WorkspaceService


def target_config(target_id: str = "alpha", *, txt_template: str | None = None) -> TargetConfig:
    return TargetConfig(
        id=target_id,
        name=f"{target_id.title()} Target",
        alias=f"{target_id.upper()}",
        title=f"{target_id.title()} Title",
        geojson_file=f"targets/{target_id}.geojson",
        coordinate=[106.7, 10.8],
        scale=50000,
        grid=GridConfig(interval=GridInterval(minutes=1)),
        export={
            "template_pptx_file": f"templates/{target_id}.pptx",
            "txt_line_template": txt_template
            or "{slide_number}|{target_id}|{capture_date}|{time_label}",
            "placeholders": [
                {
                    "field": "map",
                    "element_id": 10,
                    "kind": PlaceholderType.MAP_IMAGE,
                    "required": True,
                }
            ],
        },
        metadata={
            "template_metadata": TemplateMetadata(
                template_pptx=f"templates/{target_id}.pptx",
                slide_index=0,
                map_frame=MapFrame(x=0, y=0, width=640, height=360),
                placeholders=[
                    TemplatePlaceholder(
                        field="map",
                        element_id=10,
                        kind=PlaceholderType.MAP_IMAGE,
                        required=True,
                    )
                ],
            ).model_dump(mode="json")
        },
    )


def composition(
    composition_id: str,
    target_id: str = "alpha",
    *,
    capture_date: date = date(2026, 5, 25),
    review_order: int | None = 1,
    final_render_path: str | None = "renders/final/alpha.png",
    ready: bool = True,
    include: bool = True,
    needs_revalidation: bool = False,
) -> Composition:
    return Composition(
        composition_id=composition_id,
        target_id=target_id,
        capture_date=capture_date,
        view=ViewState(center=[106.7, 10.8], scale=50000),
        reviewed=True,
        ready=ready,
        include=include,
        needs_revalidation=needs_revalidation,
        review_order=review_order,
        artifacts=CompositionArtifacts(final_render_path=final_render_path),
        layers=[
            ImageLayer(
                layer_id=f"{composition_id}-layer",
                source_path=f"{composition_id}.tif",
                order=0,
                visible=True,
                capture_date=capture_date,
                capture_time=time(8, 30),
                metadata_status=MetadataStatus.VALID,
            )
        ],
    )


def prepare_workspace(tmp_path: Path, *compositions: Composition) -> WorkspaceService:
    service = WorkspaceService(tmp_path / "workspace")
    service.initialize(config_path="config.json")
    render_dir = service.paths.root / "renders" / "final"
    render_dir.mkdir(parents=True, exist_ok=True)
    (render_dir / "alpha.png").write_bytes(b"png")
    (render_dir / "beta.png").write_bytes(b"png")
    for item in compositions:
        service.write_composition(item)
    return service


def success_render(spec: RenderSpec, is_cancelled=None) -> RasterRenderResult:
    return RasterRenderResult(
        canvas=np.zeros((spec.output_height, spec.output_width, 3), dtype=np.uint8),
        painted_layer_ids=tuple(layer.layer_id for layer in spec.visible_layers),
    )


def issue_ids(plan) -> set[str]:  # noqa: ANN001
    return {issue.issue_id for issue in plan.issues}


def test_export_preflight_plan_sorts_by_review_order_and_summarizes(tmp_path: Path) -> None:
    beta = target_config("beta")
    alpha = target_config("alpha")
    service = prepare_workspace(
        tmp_path,
        composition(
            "beta__20260525",
            "beta",
            review_order=2,
            final_render_path="renders/final/beta.png",
        ),
        composition("alpha__20260524", "alpha", capture_date=date(2026, 5, 24), review_order=1),
    )
    ensure_final_renders_for_export(service, [beta, alpha], render=success_render)

    plan = build_export_preflight_plan(
        service,
        [beta, alpha],
    )

    assert [row.composition_id for row in plan.rows] == ["alpha__20260524", "beta__20260525"]
    assert [row.slide_number for row in plan.rows] == [1, 2]
    assert plan.summary.included_slide_count == 2
    assert plan.summary.target_count == 2
    assert plan.summary.error_count == 0


def test_export_preflight_plan_recomputes_stale_and_render_issues(tmp_path: Path) -> None:
    service = prepare_workspace(
        tmp_path,
        composition(
            "alpha__20260525",
            needs_revalidation=True,
            final_render_path=None,
        ),
    )

    plan = build_export_preflight_plan(service, [target_config("alpha")])

    assert {"composition.needs_revalidation", "export.final_render_missing"} <= issue_ids(plan)
    assert plan.summary.state.value == "blocked"
    assert plan.summary.skipped_count == 1
    assert plan.rows[0].blocking is True


def test_export_preflight_plan_surfaces_template_and_txt_issues(tmp_path: Path) -> None:
    service = prepare_workspace(tmp_path, composition("alpha__20260525"))
    warning = Issue(
        issue_id="target.template_compatibility_unknown",
        severity=IssueSeverity.WARNING,
        scope=IssueScope.TEMPLATE,
        target_id="alpha",
        message="Nhieu PPTX template chua xac minh tuong thich.",
        remediation="Kiem tra base/theme/master truoc khi export.",
    )
    target = target_config("alpha", txt_template="{unknown_token}")
    target.export.placeholders.clear()

    plan = build_export_preflight_plan(service, [target], template_issues=[warning])

    assert "target.template_compatibility_unknown" in issue_ids(plan)
    assert "export.map_placeholder_missing" in issue_ids(plan)
    assert "export.txt_placeholder_unknown" in issue_ids(plan)
    assert plan.rows[0].template_status == "ERROR"
    assert "target.template_compatibility_unknown" in {
        issue.issue_id for issue in plan.rows[0].issues
    }
    assert plan.summary.warning_count == 1


def test_export_preflight_blocks_mixed_template_slide_sizes(tmp_path: Path) -> None:
    alpha_template = tmp_path / "templates" / "alpha.pptx"
    beta_template = tmp_path / "templates" / "beta.pptx"
    _write_pptx_template(alpha_template, width=10, height=5.625)
    _write_pptx_template(beta_template, width=11, height=6.5)
    alpha = target_config("alpha")
    beta = target_config("beta")
    _set_template_path(alpha, alpha_template)
    _set_template_path(beta, beta_template)
    service = prepare_workspace(
        tmp_path,
        composition("alpha__20260525", "alpha", review_order=1),
        composition(
            "beta__20260525",
            "beta",
            review_order=2,
            final_render_path="renders/final/beta.png",
        ),
    )
    ensure_final_renders_for_export(service, [alpha, beta], render=success_render)

    plan = build_export_preflight_plan(service, [alpha, beta])

    assert "export.pptx_template_slide_size_mismatch" in issue_ids(plan)
    assert plan.summary.state.value == "blocked"
    assert plan.summary.error_count == 1
    beta_row = next(row for row in plan.rows if row.target_id == "beta")
    assert beta_row.blocking is True
    assert beta_row.template_status == "ERROR"


def _write_pptx_template(path: Path, *, width: float, height: float) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.75), Inches(4), Inches(3))
    presentation.save(path)


def _set_template_path(target: TargetConfig, template_path: Path) -> None:
    target.export.template_pptx_file = str(template_path)
    target.metadata["template_metadata"]["template_pptx"] = str(template_path)
